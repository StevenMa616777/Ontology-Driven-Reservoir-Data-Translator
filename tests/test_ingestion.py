import json

import pytest
from openpyxl import Workbook
from pydantic import ValidationError

from reservoir_data_translator.ingestion import (
    CsvParser,
    ExcelParser,
    IngestionError,
    JsonParser,
    RawBlock,
    RawDocument,
    TextParser,
    parse_document,
)


def test_raw_models_enforce_block_shape_and_unique_ids() -> None:
    with pytest.raises(ValidationError, match="row width"):
        RawBlock(
            block_id="block_0001",
            block_type="table",
            content={"columns": ["A", "B"], "rows": [[1]]},
        )

    block = RawBlock(
        block_id="block_0001",
        block_type="text",
        content="source",
    )
    with pytest.raises(ValidationError, match="must be unique"):
        RawDocument(
            source_id="demo",
            source_type="txt",
            file_name="demo.txt",
            blocks=[block, block],
        )


def test_text_parser_preserves_paragraphs_and_line_locations(tmp_path) -> None:
    path = tmp_path / "demo.txt"
    path.write_text("A15 liquid rate 500 m3/day\nminimum BHP 80 bar\n\n5 years\n")

    document = TextParser().parse(path, source_id="client-a")

    assert document.source_id == "client-a"
    assert document.source_type == "txt"
    assert [block.content for block in document.blocks] == [
        "A15 liquid rate 500 m3/day\nminimum BHP 80 bar",
        "5 years",
    ]
    assert [block.source_location for block in document.blocks] == [
        "lines 1-2",
        "line 4",
    ]


def test_json_parser_preserves_leaf_paths_and_record_tables(tmp_path) -> None:
    path = tmp_path / "client.json"
    path.write_text(
        json.dumps(
            {
                "case": {"duration": 5, "unit": "year"},
                "wells": [
                    {"Well_ID": "A15", "LiquidRate": 500},
                    {"Well_ID": "B2", "LiquidRate": 500},
                ],
            }
        )
    )

    document = JsonParser().parse(path)

    assert document.source_type == "json"
    assert document.blocks[0].content == {
        "columns": ["duration", "unit"],
        "rows": [[5, "year"]],
    }
    assert document.blocks[0].source_location == "$.case"
    assert document.blocks[1].block_type == "table"
    assert document.blocks[1].content == {
        "columns": ["Well_ID", "LiquidRate"],
        "rows": [["A15", 500], ["B2", 500]],
    }
    assert document.blocks[1].source_location == "$.wells"


def test_csv_parser_keeps_cells_as_source_strings(tmp_path) -> None:
    path = tmp_path / "wells.csv"
    path.write_text("Well,Rate,Min BHP\nA15,500,80\n", encoding="utf-8")

    document = CsvParser().parse(path)

    assert document.blocks[0].block_type == "table"
    assert document.blocks[0].content == {
        "columns": ["Well", "Rate", "Min BHP"],
        "rows": [["A15", "500", "80"]],
    }
    assert document.blocks[0].source_location == "rows 1-2"


def test_csv_parser_rejects_rows_that_cannot_form_one_table(tmp_path) -> None:
    path = tmp_path / "broken.csv"
    path.write_text("A,B\n1\n", encoding="utf-8")

    with pytest.raises(IngestionError) as error:
        CsvParser().parse(path)

    assert error.value.code == "INCONSISTENT_TABLE_WIDTH"


def test_excel_parser_emits_one_table_per_nonempty_sheet(tmp_path) -> None:
    path = tmp_path / "client.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Well Data"
    sheet.append(["Well", "LiquidRate", "MinBHP"])
    sheet.append(["A15", 500, 80])
    workbook.create_sheet("Empty")
    workbook.save(path)

    document = ExcelParser().parse(path, source_id="xlsx-client")

    assert document.source_id == "xlsx-client"
    assert len(document.blocks) == 1
    assert document.blocks[0].content == {
        "columns": ["Well", "LiquidRate", "MinBHP"],
        "rows": [["A15", 500, 80]],
    }
    assert document.blocks[0].source_location == "'Well Data'!A1:C2"


def test_parse_document_routes_supported_suffix_and_rejects_unknown(tmp_path) -> None:
    source = tmp_path / "source.TXT"
    source.write_text("oil density", encoding="utf-8")
    assert parse_document(source).source_type == "txt"

    unknown = tmp_path / "source.pdf"
    unknown.write_bytes(b"not a PDF")
    with pytest.raises(IngestionError) as error:
        parse_document(unknown)
    assert error.value.code == "UNSUPPORTED_SOURCE_TYPE"
